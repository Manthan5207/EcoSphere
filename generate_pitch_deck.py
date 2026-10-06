import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Color Palette
    BG_COLOR = RGBColor(5, 12, 14)          # #050C0E
    CARD_BG = RGBColor(13, 27, 30)          # #0D1B1E
    CARD_BORDER = RGBColor(16, 185, 129)    # #10B981
    TEXT_WHITE = RGBColor(255, 255, 255)
    TEXT_MUTED = RGBColor(148, 163, 184)    # #94A3B8
    EMERALD = RGBColor(16, 185, 129)        # #10B981
    SKY = RGBColor(56, 189, 248)            # #38BDF8
    AMBER = RGBColor(245, 158, 11)          # #F59E0B
    ROSE = RGBColor(239, 68, 68)            # #EF4444
    PURPLE = RGBColor(168, 85, 247)         # #A855F7

    def set_slide_background(slide):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = BG_COLOR

    def add_card(slide, left, top, width, height, bg_color=CARD_BG, border_color=None):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        if border_color:
            shape.line.color.rgb = border_color
            shape.line.width = Pt(1.2)
        else:
            shape.line.fill.background()
        return shape

    def add_header(slide, tag_text, title_text, subtitle_text=None, tag_color=EMERALD):
        # Tag pill
        tag_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.45), Inches(8), Inches(0.4))
        tf_tag = tag_box.text_frame
        tf_tag.word_wrap = True
        p_tag = tf_tag.paragraphs[0]
        p_tag.text = f"• {tag_text.upper()}"
        p_tag.font.size = Pt(11)
        p_tag.font.bold = True
        p_tag.font.color.rgb = tag_color
        p_tag.font.name = "Arial"

        # Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(11.7), Inches(0.8))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(28)
        p_title.font.bold = True
        p_title.font.color.rgb = TEXT_WHITE
        p_title.font.name = "Arial"

        # Subtitle
        if subtitle_text:
            p_sub = tf_title.add_paragraph()
            p_sub.text = subtitle_text
            p_sub.font.size = Pt(13)
            p_sub.font.bold = False
            p_sub.font.color.rgb = TEXT_MUTED
            p_sub.font.name = "Arial"
            p_sub.space_before = Pt(4)

    # ==========================================
    # SLIDE 1: Title & Hero
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1)
    
    # Tag
    tag1 = s1.shapes.add_textbox(Inches(1.0), Inches(1.0), Inches(5), Inches(0.4))
    p = tag1.text_frame.paragraphs[0]
    p.text = "• SMART INNOVATION PITCH"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = EMERALD

    # Main Title
    t1 = s1.shapes.add_textbox(Inches(1.0), Inches(1.5), Inches(8.5), Inches(2.2))
    tf1 = t1.text_frame
    tf1.word_wrap = True
    p1 = tf1.paragraphs[0]
    p1.text = "EcoSphere"
    p1.font.size = Pt(56)
    p1.font.bold = True
    p1.font.color.rgb = TEXT_WHITE

    p2 = tf1.add_paragraph()
    p2.text = "Understand. Simulate. Act."
    p2.font.size = Pt(24)
    p2.font.bold = True
    p2.font.color.rgb = EMERALD
    p2.space_before = Pt(6)

    p3 = tf1.add_paragraph()
    p3.text = "A unified environmental intelligence platform connecting live air quality, microclimate, personal carbon footprint, and predictive scenario simulations."
    p3.font.size = Pt(14)
    p3.font.color.rgb = TEXT_MUTED
    p3.space_before = Pt(12)

    # 4 Status Badges
    badges = [
        ("AQI INDEX", "142", "(Live Demo)", AMBER),
        ("AMBIENT TEMP", "29°C", "(Telemetry)", SKY),
        ("ECO SCORE", "78 / 100", "(Top 22%)", EMERALD),
        ("CO₂ FOOTPRINT", "1.84 T/yr", "(IPCC Calibrated)", PURPLE),
    ]

    card_w = Inches(2.7)
    card_h = Inches(1.7)
    start_x = Inches(1.0)
    start_y = Inches(4.7)

    for i, (label, val, sub, col) in enumerate(badges):
        x = start_x + i * Inches(2.85)
        add_card(s1, x, start_y, card_w, card_h, CARD_BG, col)
        
        tb = s1.shapes.add_textbox(x + Inches(0.2), start_y + Inches(0.15), card_w - Inches(0.4), card_h - Inches(0.3))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = label
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = TEXT_MUTED
        
        p_val = tf.add_paragraph()
        p_val.text = val
        p_val.font.size = Pt(22)
        p_val.font.bold = True
        p_val.font.color.rgb = col
        p_val.space_before = Pt(4)

        p_sub = tf.add_paragraph()
        p_sub.text = sub
        p_sub.font.size = Pt(10)
        p_sub.font.color.rgb = TEXT_MUTED
        p_sub.space_before = Pt(2)

    # ==========================================
    # SLIDE 2: Problem Definition
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2)
    add_header(s2, "Problem Definition", "Environmental Data Is Fragmented", "People can access environmental data — but rarely see how their choices connect to the surrounding ecosystem.", ROSE)

    # Left: 4 Pain Points
    pain_points = [
        ("AQI Apps", "Passive numbers (PM2.5, PM10) without contextual behavior recommendations.", AMBER),
        ("Weather Apps", "Isolated temperature & forecasts siloed from pollution and heat vulnerability.", SKY),
        ("Carbon Tools", "Generic annual calculators detached from real-time city conditions and daily habits.", EMERALD),
        ("Green Advice", "Vague, non-quantified lifestyle tips lacking individual attribution and accountability.", PURPLE),
    ]

    for i, (title, desc, col) in enumerate(pain_points):
        row = i // 2
        col_idx = i % 2
        px = Inches(0.8) + col_idx * Inches(2.85)
        py = Inches(2.0) + row * Inches(1.8)
        
        add_card(s2, px, py, Inches(2.7), Inches(1.65), CARD_BG, col)
        tb = s2.shapes.add_textbox(px + Inches(0.15), py + Inches(0.15), Inches(2.4), Inches(1.35))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = TEXT_WHITE
        
        p_desc = tf.add_paragraph()
        p_desc.text = desc
        p_desc.font.size = Pt(11)
        p_desc.font.color.rgb = TEXT_MUTED
        p_desc.space_before = Pt(6)

    # Right: Environmental Funnel Card
    fx = Inches(6.8)
    fy = Inches(2.0)
    fw = Inches(5.7)
    fh = Inches(4.8)
    add_card(s2, fx, fy, fw, fh, CARD_BG, ROSE)

    ftb = s2.shapes.add_textbox(fx + Inches(0.3), fy + Inches(0.2), fw - Inches(0.6), fh - Inches(0.4))
    ftf = ftb.text_frame
    ftf.word_wrap = True

    fp1 = ftf.paragraphs[0]
    fp1.text = "ENVIRONMENTAL CONVERSION FUNNEL (CRITICAL DROP-OFF)"
    fp1.font.size = Pt(12)
    fp1.font.bold = True
    fp1.font.color.rgb = ROSE

    funnel_steps = [
        ("01. INFORMATION (Access to AQI & News)", "100%", SKY),
        ("02. AWARENESS (General Concern)", "78%", SKY),
        ("📉 STEEPEST DROP (-52%): Disconnection from tangible options", "", ROSE),
        ("03. UNDERSTANDING (Local Impact)", "26%", AMBER),
        ("04. ACTION (Measurable Shift)", "9%", ROSE),
    ]

    for label, pct, col in funnel_steps:
        p = ftf.add_paragraph()
        p.text = f"{label}  {pct}"
        p.font.size = Pt(11)
        p.font.bold = True if pct or "STEEPEST" in label else False
        p.font.color.rgb = col
        p.space_before = Pt(8)

    p_res = ftf.add_paragraph()
    p_res.text = "\nThe Result: Cognitive Overload & Inaction\nUsers see high smog or high carbon stats, but cannot pinpoint which personal decision moves the needle."
    p_res.font.size = Pt(11)
    p_res.font.color.rgb = TEXT_WHITE
    p_res.space_before = Pt(10)

    # ==========================================
    # SLIDE 3: Platform Architecture
    # ==========================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3)
    add_header(s3, "Platform Architecture", "One Platform. Three Steps.", "EcoSphere links real-time environmental context to predictive consequences and personal accountability.", EMERALD)

    steps = [
        ("01", "UNDERSTAND", "Real-time hyperlocal intelligence across multi-source atmospheric sensors.",
         ["Live station & model-based AQI", "6 major criteria pollutants", "Ambient weather & humidity", "7-Day predictive forecast"],
         "Telemetry Layer: Zero latency", EMERALD),
        ("02", "SIMULATE", "Interactive digital twin modeling systemic urban and policy interventions.",
         ["What-if scenario testing", "Canopy & deforestation impacts", "+10 / +50 year horizon view", "Simplified parametric model"],
         "Predictive Layer: Decision preview", SKY),
        ("03", "ACT", "Pragmatic behavior modification with quantified carbon and cost incentives.",
         ["Dynamic personal carbon tally", "Ranked high-leverage swaps", "Eco Score tracking (0–100)", "Pledges & verified impact"],
         "Behavior Layer: Quantified habits", AMBER),
    ]

    for i, (num, title, desc, bullets, layer, col) in enumerate(steps):
        sx = Inches(0.8) + i * Inches(3.95)
        sy = Inches(2.0)
        sw = Inches(3.75)
        sh = Inches(4.5)
        
        add_card(s3, sx, sy, sw, sh, CARD_BG, col)
        tb = s3.shapes.add_textbox(sx + Inches(0.2), sy + Inches(0.2), sw - Inches(0.4), sh - Inches(0.4))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = f"{num}  {title}"
        p.font.size = Pt(20)
        p.font.bold = True
        p.font.color.rgb = col
        
        p_desc = tf.add_paragraph()
        p_desc.text = desc
        p_desc.font.size = Pt(11)
        p_desc.font.color.rgb = TEXT_MUTED
        p_desc.space_before = Pt(4)
        
        for bullet in bullets:
            pb = tf.add_paragraph()
            pb.text = f"✓ {bullet}"
            pb.font.size = Pt(11)
            pb.font.color.rgb = TEXT_WHITE
            pb.space_before = Pt(4)

        p_layer = tf.add_paragraph()
        p_layer.text = f"\n{layer}"
        p_layer.font.size = Pt(11)
        p_layer.font.bold = True
        p_layer.font.color.rgb = col
        p_layer.space_before = Pt(6)

    # Bottom Loop
    loop_box = s3.shapes.add_textbox(Inches(0.8), Inches(6.65), Inches(11.7), Inches(0.5))
    lp = loop_box.text_frame.paragraphs[0]
    lp.text = "🔄 Continuous Feedback Loop: Sense Local Air  →  Simulate Community Trajectory  →  Optimize Personal Decarbonization"
    lp.font.size = Pt(11)
    lp.font.bold = True
    lp.font.color.rgb = EMERALD
    lp.alignment = PP_ALIGN.CENTER

    # ==========================================
    # SLIDE 4: Digital Twin Engine
    # ==========================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4)
    add_header(s4, "Digital Twin Engine", "The What-If Simulator", "See the consequences before you act — interactive environmental scenario projection.", SKY)

    # Left: Baseline vs Simulated
    lx = Inches(0.8)
    ly = Inches(2.0)
    lw = Inches(5.6)
    
    # Baseline
    add_card(s4, lx, ly, lw, Inches(1.5), CARD_BG, SKY)
    tb = s4.shapes.add_textbox(lx + Inches(0.2), ly + Inches(0.15), lw - Inches(0.4), Inches(1.2))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "CURRENT CITY BASELINE (Status Quo)"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = SKY
    
    p2 = tf.add_paragraph()
    p2.text = "🌲 Tree Cover: 32%     |     🌡️ Microclimate: 29°C\n💨 Particulate Smog: AQI 85     |     🏥 Health Risk: Baseline"
    p2.font.size = Pt(11)
    p2.font.color.rgb = TEXT_WHITE
    p2.space_before = Pt(4)

    # Simulated
    add_card(s4, lx, ly + Inches(1.7), lw, Inches(1.5), CARD_BG, ROSE)
    tb = s4.shapes.add_textbox(lx + Inches(0.2), ly + Inches(1.85), lw - Inches(0.4), Inches(1.2))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "SIMULATED SCENARIO (+50 YEARS - Severe Impact)"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = ROSE
    
    p2 = tf.add_paragraph()
    p2.text = "⚠️ Tree Cover: 11% (-21%)     |     🔥 Heat Island: +3.8°C\n🌫️ Particulate Smog: AQI 168     |     🫁 Respiratory Burden: +42%"
    p2.font.size = Pt(11)
    p2.font.color.rgb = TEXT_WHITE
    p2.space_before = Pt(4)

    # Insight
    add_card(s4, lx, ly + Inches(3.4), lw, Inches(1.4), CARD_BG, AMBER)
    tb = s4.shapes.add_textbox(lx + Inches(0.2), ly + Inches(3.5), lw - Inches(0.4), Inches(1.2))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Key Scenario Insight:"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = AMBER
    p2 = tf.add_paragraph()
    p2.text = "Urban tree loss triggers compounding feedback loops — surface heat absorption surges by 26%, doubling stagnation of PM2.5 particulates over 50 years."
    p2.font.size = Pt(10.5)
    p2.font.color.rgb = TEXT_WHITE
    p2.space_before = Pt(2)

    # Right: Trajectory Trajectory Box
    rx = Inches(6.7)
    rw = Inches(5.8)
    rh = Inches(4.8)
    add_card(s4, rx, ly, rw, rh, CARD_BG, EMERALD)
    tb = s4.shapes.add_textbox(rx + Inches(0.2), ly + Inches(0.2), rw - Inches(0.4), rh - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = "PARAMETRIC SCENARIO TRAJECTORY"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = TEXT_WHITE

    p_leg = tf.add_paragraph()
    p_leg.text = "● Temp & Smog Index (Red)   |   ● Green Resilience (Green)"
    p_leg.font.size = Pt(10)
    p_leg.font.color.rgb = TEXT_MUTED
    p_leg.space_before = Pt(4)

    traj_lines = [
        ("\n1. Current (2026 - Status Quo):", "Temp & Smog: Low (AQI 85, 29°C)  |  Green Resilience: 70%", EMERALD),
        ("\n2. +10 Years Horizon (2036):", "Temp & Smog: Rising (+18 pts)  |  Green Resilience: 48%", AMBER),
        ("\n3. +50 Years Horizon (2076):", "Temp & Smog: Severe (AQI 168, +3.8°C)  |  Green Resilience: 18%", ROSE),
    ]

    for title_t, desc_t, col in traj_lines:
        pt = tf.add_paragraph()
        pt.text = title_t
        pt.font.size = Pt(11)
        pt.font.bold = True
        pt.font.color.rgb = col
        
        pd = tf.add_paragraph()
        pd.text = desc_t
        pd.font.size = Pt(10)
        pd.font.color.rgb = TEXT_WHITE

    # ==========================================
    # SLIDE 5: Unified Environmental Dashboard
    # ==========================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5)
    add_header(s5, "Feature Suite", "Unified Environmental Dashboard", "Everything in one environmental cockpit — live sensors, carbon diagnostics, and predictive insights.", EMERALD)

    # 3 Top Cockpit Cards
    cockpit_cards = [
        ("AIR QUALITY INDEX", "142", "Unhealthy", "PM2.5 Primary Driver\nSensitive groups should limit outdoor activity.", ROSE),
        ("MICROCLIMATE TELEMETRY", "29°C", "New Delhi", "Partly Cloudy\n💧 Humidity: 64%  •  💨 Wind: 14 km/h NW", SKY),
        ("PERSONAL ECO SCORE", "78 / 100", "Top 22%", "+6 pts this month\nAnnual Footprint: 1,840 kg CO₂/yr", EMERALD),
    ]

    for i, (head, val, tag, details, col) in enumerate(cockpit_cards):
        cx = Inches(0.8) + i * Inches(3.95)
        cy = Inches(2.0)
        cw = Inches(3.75)
        ch = Inches(1.8)
        
        add_card(s5, cx, cy, cw, ch, CARD_BG, col)
        tb = s5.shapes.add_textbox(cx + Inches(0.2), cy + Inches(0.15), cw - Inches(0.4), ch - Inches(0.3))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = f"{head}  [{tag}]"
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = col
        
        pv = tf.add_paragraph()
        pv.text = val
        pv.font.size = Pt(24)
        pv.font.bold = True
        pv.font.color.rgb = TEXT_WHITE
        pv.space_before = Pt(2)
        
        pdet = tf.add_paragraph()
        pdet.text = details
        pdet.font.size = Pt(10)
        pdet.font.color.rgb = TEXT_MUTED
        pdet.space_before = Pt(2)

    # Bottom Row: 7-Day Forecast & 6-Pollutants Matrix
    add_card(s5, Inches(0.8), Inches(4.0), Inches(5.75), Inches(2.0), CARD_BG, SKY)
    tb = s5.shapes.add_textbox(Inches(1.0), Inches(4.1), Inches(5.35), Inches(1.8))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "7-DAY ENVIRONMENTAL FORECAST"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = SKY
    p_fc = tf.add_paragraph()
    p_fc.text = "Mon: 29°C (AQI 138)  |  Tue: 30°C (AQI 142)  |  Wed: 31°C (AQI 145)\nThu: 29°C (AQI 135)  |  Fri: 28°C (AQI 120)  |  Sat: 30°C (AQI 130)\nSun: 29°C (AQI 115)"
    p_fc.font.size = Pt(10.5)
    p_fc.font.color.rgb = TEXT_WHITE
    p_fc.space_before = Pt(4)

    add_card(s5, Inches(6.75), Inches(4.0), Inches(5.75), Inches(2.0), CARD_BG, EMERALD)
    tb = s5.shapes.add_textbox(Inches(6.95), Inches(4.1), Inches(5.35), Inches(1.8))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "6-CATEGORY POLLUTANT MATRIX (CONCENTRATION)"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = EMERALD
    p_pol = tf.add_paragraph()
    p_pol.text = "• PM2.5 (Fine Particulates): 84% (High)   |   • PM10 (Coarse Dust): 62%\n• NO₂ (Vehicle Exhaust): 52%             |   • SO₂ (Industrial Gas): 24%\n• CO (Carbon Monoxide): 31%             |   • O₃ (Ground Ozone): 41%"
    p_pol.font.size = Pt(10.5)
    p_pol.font.color.rgb = TEXT_WHITE
    p_pol.space_before = Pt(4)

    # Top Action Bar
    add_card(s5, Inches(0.8), Inches(6.15), Inches(11.7), Inches(0.75), CARD_BG, EMERALD)
    tb = s5.shapes.add_textbox(Inches(1.0), Inches(6.2), Inches(11.3), Inches(0.65))
    p = tb.text_frame.paragraphs[0]
    p.text = "⭐ TOP ACTION: Switch 2 commute days/week to electric metro: Potential saving of 320 kg CO₂/year  (+8 Eco Score Points)"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = EMERALD

    # ==========================================
    # SLIDE 6: Behavioral Engine
    # ==========================================
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_background(s6)
    add_header(s6, "Behavioral Engine", "From Numbers to Decisions", "Translating abstract carbon metrics into ranked, achievable lifestyle adjustments.", EMERALD)

    # 4 Steps Pill
    pipeline = [
        ("1. Lifestyle Inputs", "Transport, Energy, Diet", SKY),
        ("2. Footprint Score", "1,840 kg CO₂/yr", EMERALD),
        ("3. Benchmarks", "India 1.9T | Global 4.7T", AMBER),
        ("4. Targeted Actions", "Ranked Interventions", PURPLE)
    ]
    for i, (title_p, val_p, col) in enumerate(pipeline):
        px = Inches(0.8) + i * Inches(2.95)
        add_card(s6, px, Inches(2.0), Inches(2.8), Inches(0.9), CARD_BG, col)
        tb = s6.shapes.add_textbox(px + Inches(0.1), Inches(2.05), Inches(2.6), Inches(0.8))
        tf = tb.text_frame
        p = tf.paragraphs[0]
        p.text = title_p
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = col
        pv = tf.add_paragraph()
        pv.text = val_p
        pv.font.size = Pt(10.5)
        pv.font.bold = True
        pv.font.color.rgb = TEXT_WHITE

    # Left Ranked Interventions
    add_card(s6, Inches(0.8), Inches(3.1), Inches(6.8), Inches(3.8), CARD_BG, EMERALD)
    tb = s6.shapes.add_textbox(Inches(1.0), Inches(3.25), Inches(6.4), Inches(3.5))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "POTENTIAL ANNUAL CO₂ SAVINGS (RANKED)"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = EMERALD

    savings = [
        ("1. Commute Mode Shift (Metro/Bus 3 days/week)", "420 kg CO₂/yr", EMERALD),
        ("2. Home Energy Optimization (BLDC fans, LED + AC 24°C)", "280 kg CO₂/yr", SKY),
        ("3. Plant-Rich Diet Swaps (4 low-footprint meals/week)", "190 kg CO₂/yr", AMBER),
        ("4. Consumption & Longevity (Reduced fast fashion)", "110 kg CO₂/yr", PURPLE),
    ]
    for act, kg, col in savings:
        p_act = tf.add_paragraph()
        p_act.text = f"{act}  →  {kg}"
        p_act.font.size = Pt(11)
        p_act.font.bold = True
        p_act.font.color.rgb = col
        p_act.space_before = Pt(8)

    # Right Tree Translation & Rupee Dividend
    add_card(s6, Inches(7.8), Inches(3.1), Inches(4.7), Inches(1.8), CARD_BG, EMERALD)
    tb = s6.shapes.add_textbox(Inches(8.0), Inches(3.2), Inches(4.3), Inches(1.6))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "TREE-EQUIVALENT TRANSLATION\n≈ 20 Mature Trees"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = EMERALD
    p2 = tf.add_paragraph()
    p2.text = "A 420 kg CO₂ reduction equals the annual carbon sequestered by ~20 growing urban trees."
    p2.font.size = Pt(10)
    p2.font.color.rgb = TEXT_WHITE
    p2.space_before = Pt(4)

    add_card(s6, Inches(7.8), Inches(5.1), Inches(2.25), Inches(1.8), CARD_BG, SKY)
    tb = s6.shapes.add_textbox(Inches(7.9), Inches(5.2), Inches(2.05), Inches(1.6))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "ANNUAL COST DIVIDEND\n\n₹14,200\nSaved / Year"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = SKY

    add_card(s6, Inches(10.25), Inches(5.1), Inches(2.25), Inches(1.8), CARD_BG, AMBER)
    tb = s6.shapes.add_textbox(Inches(10.35), Inches(5.2), Inches(2.05), Inches(1.6))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "COMMUNITY BADGE\n\n🏅 Delhi Transit Star\n(Top 15% Commute)"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = AMBER

    # ==========================================
    # SLIDE 7: Engineering Stack
    # ==========================================
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_background(s7)
    add_header(s7, "Engineering Stack", "Lightweight & Scalable Tech", "Built for rapid response, zero API key dependencies, and seamless mobile execution.", SKY)

    layers = [
        ("01. CLIENT TIER", "Mobile Browser  •  Desktop Web App  •  PWA Capability", EMERALD),
        ("02. REACT CORE", "React 18 + Vite  •  Tailwind CSS  •  Framer Motion  •  Recharts SVG Engine  •  Zustand Store", SKY),
        ("03. APP ENGINES", "Carbon Calculator  •  Scenario Engine (+10/+50Y)  •  Eco Score Normalizer  •  Recommendation Filter", AMBER),
        ("04. DATA LAYER", "Open-Meteo APIs (Weather)  •  Model-Based AQI Feeds  •  Offline Sample Dataset Fallback", PURPLE),
    ]

    for i, (title_l, desc_l, col) in enumerate(layers):
        ly = Inches(2.0) + i * Inches(1.05)
        add_card(s7, Inches(0.8), ly, Inches(11.7), Inches(0.9), CARD_BG, col)
        tb = s7.shapes.add_textbox(Inches(1.0), ly + Inches(0.15), Inches(11.3), Inches(0.6))
        tf = tb.text_frame
        p = tf.paragraphs[0]
        p.text = f"{title_l}  —  {desc_l}"
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = col

    # Bottom 4 pillars
    stack_pills = [
        ("⚡ Sub-Second Speed", "Vite optimized SPA bundle"),
        ("📱 Mobile-First Design", "Tailwind fluid viewport"),
        ("📦 Offline Fallback", "Preloaded demo caches"),
        ("🔑 Zero API Key Cost", "Open-Meteo open tier")
    ]
    for i, (h, sub) in enumerate(stack_pills):
        px = Inches(0.8) + i * Inches(2.95)
        add_card(s7, px, Inches(6.3), Inches(2.8), Inches(0.8), CARD_BG, EMERALD)
        tb = s7.shapes.add_textbox(px + Inches(0.1), Inches(6.35), Inches(2.6), Inches(0.7))
        tf = tb.text_frame
        p = tf.paragraphs[0]
        p.text = h
        p.font.size = Pt(10.5)
        p.font.bold = True
        p.font.color.rgb = EMERALD
        pv = tf.add_paragraph()
        pv.text = sub
        pv.font.size = Pt(9)
        pv.font.color.rgb = TEXT_MUTED

    # ==========================================
    # SLIDE 8: Growth & Roadmap
    # ==========================================
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_background(s8)
    add_header(s8, "Growth & Roadmap", "Making Awareness Actionable", "Bridging individual actions with city-scale climate resilience.", EMERALD)

    # 4 Impact Pillars
    add_card(s8, Inches(0.8), Inches(2.0), Inches(5.6), Inches(4.5), CARD_BG, SKY)
    tb = s8.shapes.add_textbox(Inches(1.0), Inches(2.15), Inches(5.2), Inches(4.2))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "CORE IMPACT PILLARS (4 STRATEGIC VECTORS)"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = SKY

    pillars = [
        ("👤 Personalized", "Translates macro gigatons into direct personal kilograms and household savings."),
        ("📊 Visual", "Replaces complex sensor tables with intuitive scenarios and progress telemetry."),
        ("📍 India-Focused", "Tailored for Delhi-NCR microclimates, INR currency savings, and transit routes."),
        ("🎓 Expandable", "Ready for campus hackathons, colleges, schools, and civic green leagues.")
    ]
    for pil, desc in pillars:
        p1 = tf.add_paragraph()
        p1.text = pil
        p1.font.size = Pt(11)
        p1.font.bold = True
        p1.font.color.rgb = TEXT_WHITE
        p1.space_before = Pt(6)
        
        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(10)
        p2.font.color.rgb = TEXT_MUTED

    # 3 Execution Phases
    add_card(s8, Inches(6.7), Inches(2.0), Inches(5.8), Inches(4.5), CARD_BG, EMERALD)
    tb = s8.shapes.add_textbox(Inches(6.9), Inches(2.15), Inches(5.4), Inches(4.2))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "EVOLUTIONARY ROADMAP (EXECUTION PHASES)"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = EMERALD

    phases = [
        ("PHASE 1 — NOW (MVP)  [COMPLETED]", "Unified telemetry dashboard, personal carbon engine, 3-scenario simulator.", EMERALD),
        ("PHASE 2 — NEXT (Q3)  [IN PROGRESS]", "Station-level sensor APIs, user auth, college cohort leaderboards & badges.", SKY),
        ("PHASE 3 — FUTURE HORIZON  [PLANNED]", "Smart meter feeds, transit smart-card sync, regional languages & NGO credits.", PURPLE)
    ]
    for ph, desc, col in phases:
        p1 = tf.add_paragraph()
        p1.text = ph
        p1.font.size = Pt(11)
        p1.font.bold = True
        p1.font.color.rgb = col
        p1.space_before = Pt(8)
        
        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(10)
        p2.font.color.rgb = TEXT_WHITE

    p_target = tf.add_paragraph()
    p_target.text = "\n🎯 Goal: 50,000 active student users across North Indian campuses."
    p_target.font.size = Pt(11)
    p_target.font.bold = True
    p_target.font.color.rgb = EMERALD
    p_target.space_before = Pt(10)

    # Bottom Vision
    v_box = s8.shapes.add_textbox(Inches(0.8), Inches(6.65), Inches(11.7), Inches(0.5))
    vp = v_box.text_frame.paragraphs[0]
    vp.text = 'OUR VISION: "Build a personal environmental intelligence layer for everyday planetary decisions."'
    vp.font.size = Pt(11)
    vp.font.bold = True
    vp.font.color.rgb = TEXT_WHITE
    vp.alignment = PP_ALIGN.CENTER

    # ==========================================
    # SLIDE 9: Questions & Discussion
    # ==========================================
    s9 = prs.slides.add_slide(blank_layout)
    set_slide_background(s9)

    tb = s9.shapes.add_textbox(Inches(1.0), Inches(1.5), Inches(11.3), Inches(4.5))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "• COLLEGE INNOVATION PROJECT PITCH"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = EMERALD
    p.alignment = PP_ALIGN.CENTER

    p1 = tf.add_paragraph()
    p1.text = "EcoSphere"
    p1.font.size = Pt(56)
    p1.font.bold = True
    p1.font.color.rgb = TEXT_WHITE
    p1.alignment = PP_ALIGN.CENTER
    p1.space_before = Pt(6)

    p2 = tf.add_paragraph()
    p2.text = "Understand. Simulate. Act."
    p2.font.size = Pt(22)
    p2.font.bold = True
    p2.font.color.rgb = EMERALD
    p2.alignment = PP_ALIGN.CENTER
    p2.space_before = Pt(4)

    p3 = tf.add_paragraph()
    p3.text = "\n\n💬 Questions & Discussion"
    p3.font.size = Pt(26)
    p3.font.bold = True
    p3.font.color.rgb = TEXT_WHITE
    p3.alignment = PP_ALIGN.CENTER

    p4 = tf.add_paragraph()
    p4.text = "Thank you for your time. We welcome feedback on our simulation engine, telemetry integration, and campus deployment roadmap."
    p4.font.size = Pt(13)
    p4.font.color.rgb = TEXT_MUTED
    p4.alignment = PP_ALIGN.CENTER
    p4.space_before = Pt(8)

    # 3 Pill badges at bottom
    pills_s9 = ["📡 Real-Time Telemetry", "🏙️ What-If Digital Twin", "🌱 Personal Decarbonization"]
    for i, pl in enumerate(pills_s9):
        px = Inches(2.2) + i * Inches(3.1)
        add_card(s9, px, Inches(6.0), Inches(2.8), Inches(0.65), CARD_BG, EMERALD)
        tb_p = s9.shapes.add_textbox(px, Inches(6.05), Inches(2.8), Inches(0.55))
        p = tb_p.text_frame.paragraphs[0]
        p.text = pl
        p.font.size = Pt(10.5)
        p.font.bold = True
        p.font.color.rgb = EMERALD
        p.alignment = PP_ALIGN.CENTER

    # ==========================================
    # SLIDE 10: Sources
    # ==========================================
    s10 = prs.slides.add_slide(blank_layout)
    set_slide_background(s10)
    add_header(s10, "Appendix & References", "Image & Data Sources", "Research calibration and imagery references for EcoSphere.", TEXT_MUTED)

    sources = [
        ("Futuristic Smart City Digital Twin Visuals",
         "https://www.ierek.com/news/wp-content/uploads/2025/06/Black-and-Green-Modern-Futuristic-Digital-Transformation-with-AI-Technology.png",
         "Source: www.ierek.com", EMERALD),
        ("Atmospheric Telemetry & Vector Assets",
         "https://elements-resized.envatousercontent.com/elements-video-cover-images/files/7f6bf28c-37b3-4559-ab34-f8c2f41de0c4/inline_image_preview.jpg",
         "Source: elements.envato.com", SKY),
        ("Scientific Carbon & Meteorological Calibration",
         "Open-Meteo Air Quality & Weather API, IPCC AR6 WG3 Climate Mitigation Benchmarks, and CEA India Baseline CO₂ Emission Database.",
         "Source: IPCC & CEA India", AMBER),
    ]

    for i, (title_s, url_s, src_s, col) in enumerate(sources):
        sy = Inches(2.0) + i * Inches(1.6)
        add_card(s10, Inches(0.8), sy, Inches(11.7), Inches(1.35), CARD_BG, col)
        tb = s10.shapes.add_textbox(Inches(1.0), sy + Inches(0.15), Inches(11.3), Inches(1.05))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = title_s
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = col
        
        pu = tf.add_paragraph()
        pu.text = url_s
        pu.font.size = Pt(10)
        pu.font.color.rgb = TEXT_WHITE
        pu.space_before = Pt(2)
        
        ps = tf.add_paragraph()
        ps.text = src_s
        ps.font.size = Pt(9.5)
        ps.font.color.rgb = TEXT_MUTED
        ps.space_before = Pt(2)

    output_path = "EcoSphere_Pitch_Deck.pptx"
    prs.save(output_path)
    print(f"Successfully generated {output_path} with {len(prs.slides)} slides!")

if __name__ == "__main__":
    create_deck()
